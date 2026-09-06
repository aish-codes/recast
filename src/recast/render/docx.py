"""DOCX output.

Worth having even though the PDF is prettier: a number of ATS parse Word more
reliably than PDF, and some portals only accept .doc/.docx.

Kept deliberately plain — no tables, no text boxes, no columns, built-in styles
only. Every one of those is a known parser trap.
"""

from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from docx.shared import Pt

from ..models.tailored import CoverLetter, TailoredResume


def _heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(11)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)


def resume_docx(resume: TailoredResume, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    _resume_doc(resume).save(str(path))
    return path


def resume_docx_bytes(resume: TailoredResume) -> bytes:
    """Same document, straight to memory — the API has no filesystem to use."""
    buf = io.BytesIO()
    _resume_doc(resume).save(buf)
    return buf.getvalue()


def _resume_doc(resume: TailoredResume) -> Document:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    name = doc.add_paragraph()
    run = name.add_run(resume.contact.name)
    run.bold = True
    run.font.size = Pt(18)
    name.paragraph_format.space_after = Pt(0)

    bits = [resume.contact.email, resume.contact.phone, resume.contact.location]
    bits += list(resume.contact.links.values())
    contact = doc.add_paragraph(" · ".join(b for b in bits if b))
    contact.paragraph_format.space_after = Pt(6)

    if resume.summary:
        _heading(doc, "Summary")
        doc.add_paragraph(resume.summary.text)

    if resume.experience:
        _heading(doc, "Experience")
        for e in resume.experience:
            p = doc.add_paragraph()
            p.add_run(f"{e.title}, {e.company}").bold = True
            p.add_run(f"  |  {e.start} – {e.end or 'Present'}")
            if e.location:
                p.add_run(f"  |  {e.location}")
            p.paragraph_format.space_after = Pt(0)
            for b in e.bullets:
                doc.add_paragraph(b.text, style="List Bullet")

    if resume.projects:
        _heading(doc, "Projects")
        for proj in resume.projects:
            p = doc.add_paragraph()
            p.add_run(proj.name).bold = True
            if proj.tech:
                p.add_run(f" — {', '.join(proj.tech)}")
            if proj.url:
                p.add_run(f"  |  {proj.url}")
            p.paragraph_format.space_after = Pt(0)
            for b in proj.bullets:
                doc.add_paragraph(b.text, style="List Bullet")

    if resume.skills:
        _heading(doc, "Skills")
        for g in resume.skills:
            p = doc.add_paragraph()
            p.add_run(f"{g.name}: ").bold = True
            p.add_run(", ".join(g.skills))
            p.paragraph_format.space_after = Pt(0)

    if resume.education:
        _heading(doc, "Education")
        for ed in resume.education:
            p = doc.add_paragraph()
            p.add_run(f"{ed.degree}{', ' + ed.field if ed.field else ''}, {ed.institution}").bold = True
            if ed.start or ed.end:
                p.add_run(f"  |  {ed.start or ''}{' – ' if ed.start and ed.end else ''}{ed.end or ''}")
            if ed.detail:
                doc.add_paragraph(ed.detail)

    if resume.certifications:
        _heading(doc, "Certifications")
        for c in resume.certifications:
            doc.add_paragraph(c, style="List Bullet")

    return doc


def cover_letter_docx(letter: CoverLetter, path: Path) -> Path:
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    head = doc.add_paragraph()
    head.add_run(letter.name).bold = True

    doc.add_paragraph(letter.greeting)
    for para in letter.paragraphs:
        doc.add_paragraph(para)
    doc.add_paragraph(letter.signoff)
    doc.add_paragraph(letter.name)

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path
